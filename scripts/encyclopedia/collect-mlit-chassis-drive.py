#!/usr/bin/env python3
"""Extract only existing chassis codes, checkpointing each official workbook separately."""
import json, pathlib, urllib.request, concurrent.futures, hashlib, re, unicodedata, io
import openpyxl,xlrd
import argparse
p=argparse.ArgumentParser();p.add_argument('--cache',required=True);p.add_argument('--jobs',required=True);p.add_argument('--inventory',default='reports/japan-chassis/codes-0000.json');args=p.parse_args()
ROOT=pathlib.Path(args.cache);ROOT.mkdir(parents=True,exist_ok=True)
inv=json.load(open(args.inventory))
def norm(x):return unicodedata.normalize('NFKC',str(x or '')).strip().upper()
def base(x):return re.sub(r'^[A-Z0-9]{1,3}-','',norm(x))
wanted={base(x['chassisCode']) for x in inv}
def process(job):
 page,url=job; name=url.rsplit('/',1)[1]; out=ROOT/(name+'.v3.json')
 if out.exists():return json.load(open(out))
 try:
  path=ROOT/name
  if not path.exists():path.write_bytes(urllib.request.urlopen(url,timeout=40).read())
  b=path.read_bytes(); sheets=[]
  if name.endswith('.xlsx'):
   wb=openpyxl.load_workbook(io.BytesIO(b),data_only=True)
   for ws in wb:
    rows=[list(r) for r in ws.values]
    for rg in ws.merged_cells.ranges:
     v=ws.cell(rg.min_row,rg.min_col).value
     for r in range(rg.min_row-1,rg.max_row):
      for c in range(rg.min_col-1,rg.max_col):rows[r][c]=v
    sheets.append((ws.title,rows))
  else:
   wb=xlrd.open_workbook(file_contents=b,formatting_info=True)
   for ws in wb.sheets():
    rows=[ws.row_values(r) for r in range(ws.nrows)]
    for a,z,c,d in ws.merged_cells:
     v=rows[a][c]
     for r in range(a,z):
      for k in range(c,d):rows[r][k]=v
    sheets.append((ws.name,rows))
  evidence=[]
  for sheet,rows in sheets:
   dc=set()
   for row in rows[:25]:
    for c,v in enumerate(row):
     if '駆動' in re.sub(r'\s','',str(v)):dc.add(c)
   brand_context=sheet
   for ri,row in enumerate(rows):
    if row and isinstance(row[0],str) and row[0].strip():brand_context=norm(row[0])
    ds={norm(row[c]) for c in dc if c<len(row)}&{'F','R','A'}
    if len(ds)!=1:continue
    for ci,v in enumerate(row):
     code=base(v)
     if code in wanted and re.fullmatch(r'[A-Z0-9]{1,3}-[A-Z0-9]+',norm(v)):
      evidence.append(dict(code=code,fullCode=norm(v),drive={'F':'fwd','R':'rwd','A':'awd'}[next(iter(ds))],sheet=sheet,brand=brand_context,row=ri+1,column=ci+1,identity=[norm(v) for v in row[:7]],url=url,page=page,sha256=hashlib.sha256(b).hexdigest()))
  result=dict(url=url,evidence=evidence,sheets=len(sheets));out.write_text(json.dumps(result,ensure_ascii=False));return result
 except Exception as e:return dict(url=url,error=str(e))
jobs=json.load(open(args.jobs))
# duplicate links need only one download
jobs=list({u:(p,u) for p,u in jobs}.values())
with concurrent.futures.ThreadPoolExecutor(max_workers=4) as ex:
 results=[]
 for i,r in enumerate(ex.map(process,jobs)):
  results.append(r)
  if i%20==0:print(i,len(jobs),'matches',sum(len(x.get('evidence',[])) for x in results),flush=True)
json.dump([dict(url=r['url'],error=r.get('error'),rows=len(r.get('evidence',[]))) for r in results],open(ROOT/'extraction-summary.json','w'),ensure_ascii=False)
print('DONE',len(results),'errors',sum('error' in r for r in results),'codes',len({e['code'] for r in results for e in r.get('evidence',[])}),flush=True)
